import asyncio
import io
import json
import re
import time
import zipfile
from types import SimpleNamespace
from unittest.mock import Mock

import httpx
import pytest
from azure.core.credentials import AccessToken
from azure.core.exceptions import ClientAuthenticationError
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from intent_evaluator import COLUMNS, MODE, QUESTION, ROOT, THRESHOLD, EvaluationCase, summarize
from web_demo import (
    DAILY_ROWS,
    MAX_UPLOAD,
    AzureSettings,
    Store,
    WebRunner,
    WebSettings,
    create_app,
    excel_template,
    parse_excel,
)

ORIGIN = "http://127.0.0.1:8877"
HEADERS = {"Origin": ORIGIN, "X-Demo-Request": "1"}


@pytest.fixture(autouse=True)
def azure_environment(monkeypatch):
    monkeypatch.setenv("AZURE_ENDPOINT", "https://test-resource.services.ai.azure.com")
    monkeypatch.setenv("AZURE_DEPLOYMENT", "test-deployment")


def workbook(rows, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.append(headers or ["id", "message", "predicted_intent", "expected_intent"])
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def report_for(cases):
    return {
        "mode": MODE,
        "judgment_threshold": THRESHOLD,
        "created_at": "now",
        "providers": ["microsoft", "jev"],
        "question": QUESTION,
        "dataset_sha256": "hash",
        "cases": [c.model_dump() for c in cases],
        "results": [],
        "summary": {p: summarize(cases, [], p) for p in ("microsoft", "jev")},
    }


@pytest.fixture
def web(tmp_path):
    settings = WebSettings(
        origin=ORIGIN,
        jev_key="test-secret-key",
        data_dir=tmp_path,
        secure_cookies=False,
    )

    class FakeRunner:
        def __init__(self, settings, azure, store):
            self.credential = SimpleNamespace(close=Mock())
            self.store = store
            self.started = []

        async def run(self, identifier, report):
            self.started.append(identifier)
            self.store.update(identifier, "completed", report)

    app = create_app(settings, runner_factory=FakeRunner)
    with TestClient(app, base_url=ORIGIN) as client:
        yield SimpleNamespace(client=client, app=app, settings=settings)


def initialize_session(client):
    response = client.get("/api/config")
    assert response.status_code == 200
    return response


def test_web_settings(tmp_path):
    with pytest.raises(ValueError):
        WebSettings("http://public.example", "secret", tmp_path, False)
    with pytest.raises(ValueError):
        WebSettings("https://demo.example", " ", tmp_path)
    with pytest.raises(ValueError):
        WebSettings("https://demo.example/", "secret", tmp_path)


def test_web_settings_do_not_require_password(monkeypatch, tmp_path):
    monkeypatch.delenv("DEMO_PASSWORD_HASH", raising=False)
    monkeypatch.setenv("PUBLIC_ORIGIN", "https://demo.example")
    monkeypatch.setenv("JEV_API_KEY", "test-secret-key")
    monkeypatch.setenv("DEMO_DATA_DIR", str(tmp_path))
    settings = WebSettings.from_env()
    assert settings.origin == "https://demo.example"
    assert settings.data_dir == tmp_path


def test_azure_settings_require_customer_configuration(monkeypatch):
    settings = AzureSettings.from_env()
    assert settings.endpoint == "https://test-resource.services.ai.azure.com"
    assert settings.deployment == "test-deployment"
    monkeypatch.delenv("AZURE_ENDPOINT")
    with pytest.raises(KeyError, match="AZURE_ENDPOINT"):
        AzureSettings.from_env()
    monkeypatch.setenv("AZURE_ENDPOINT", settings.endpoint)
    monkeypatch.delenv("AZURE_DEPLOYMENT")
    with pytest.raises(KeyError, match="AZURE_DEPLOYMENT"):
        AzureSettings.from_env()


@pytest.mark.parametrize(
    "endpoint",
    [
        "http://test-resource.services.ai.azure.com",
        "https://untrusted.example",
        "https://test-resource.services.ai.azure.com.untrusted.example",
        "https://user:password@test-resource.services.ai.azure.com",
        "https://test-resource.services.ai.azure.com:8443",
        "https://test-resource.services.ai.azure.com/path",
        "https://test-resource.services.ai.azure.com?query=value",
        "https://test-resource.services.ai.azure.com#fragment",
    ],
)
def test_azure_settings_reject_invalid_endpoints(endpoint):
    with pytest.raises(ValueError, match="AZURE_ENDPOINT"):
        AzureSettings(endpoint, "test-deployment")


def test_azure_settings_allow_supported_endpoints_and_reject_empty_deployment():
    for endpoint in (
        "https://test-resource.services.ai.azure.com",
        "https://test-resource.cognitiveservices.azure.com/",
    ):
        assert AzureSettings(endpoint, "deployment").endpoint == endpoint
    with pytest.raises(ValueError, match="AZURE_DEPLOYMENT"):
        AzureSettings("https://test-resource.services.ai.azure.com", " ")


def test_template_contains_only_headers():
    book = load_workbook(io.BytesIO(excel_template()))
    try:
        assert len(book.worksheets) == 1
        assert list(book.active.values) == [tuple(COLUMNS)]
    finally:
        book.close()


def test_template_round_trip_after_adding_case():
    book = load_workbook(io.BytesIO(excel_template()))
    try:
        book.active.append(["X", "Explain how this works", "get_information", "get_information"])
        stream = io.BytesIO()
        book.save(stream)
    finally:
        book.close()
    cases = parse_excel(stream.getvalue())
    assert len(cases) == 1
    assert cases[0].id == "X"
    assert cases[0].classifier_is_correct


@pytest.mark.parametrize(
    "rows",
    [
        [],
        [["X", "=1+1", "other", "other"]],
        [["X", "message", "other", "unknown"]],
        [["X", "message", "unknown", "other"]],
        [["X", 123, "other", "other"]],
        [["X", "a" * 1001, "other", "other"]],
        [["X" * 51, "message", "other", "other"]],
        [["X", "message", "other", "other"], ["X", "again", "other", "other"]],
        [[str(n), "message", "other", "other"] for n in range(101)],
    ],
)
def test_rejects_invalid_workbooks(rows):
    with pytest.raises(ValueError):
        parse_excel(workbook(rows))


def test_upload_shape_and_archive_limits():
    assert (
        len(parse_excel(workbook([[str(n), "message", "other", "other"] for n in range(100)])))
        == 100
    )
    with pytest.raises(ValueError):
        parse_excel(b"not a workbook")
    with pytest.raises(ValueError):
        parse_excel(b"x" * (MAX_UPLOAD + 1))
    with pytest.raises(ValueError):
        parse_excel(workbook([["X", "message", "other"]], ["id", "text", "expected_intent"]))
    with pytest.raises(ValueError, match="evaluator template"):
        parse_excel(workbook([["X", "message", "other"]], ["id", "message", "expected_intent"]))
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", "x" * 2_000_001)
    with pytest.raises(ValueError, match="too complex"):
        parse_excel(stream.getvalue())


def test_anonymous_cookies_and_no_secret_leaks(web):
    assert web.client.get("/").status_code == 200
    assert web.client.get("/health").json() == {"status": "ok"}
    response = initialize_session(web.client)
    cookie = response.headers["set-cookie"]
    assert "HttpOnly" in cookie and "SameSite=strict" in cookie
    assert "Max-Age=7200" in cookie
    token = web.client.cookies.get("demo_session")
    config = web.client.get("/api/config")
    assert config.status_code == 200
    assert "set-cookie" not in config.headers
    assert web.client.cookies.get("demo_session") == token
    assert config.json()["mode"] == MODE
    assert config.json()["question"]["type"] == "noul"
    assert "test-secret-key" not in config.text + response.text
    assert config.headers["cache-control"] == "no-store"
    assert "'unsafe-inline'" not in config.headers["content-security-policy"]
    template = web.client.get("/api/template")
    assert template.status_code == 200
    assert 'filename="evaluator-template.xlsx"' in template.headers["content-disposition"]
    book = load_workbook(io.BytesIO(template.content))
    try:
        assert len(book.worksheets) == 1
        assert list(book.active.values) == [tuple(COLUMNS)]
    finally:
        book.close()


@pytest.mark.parametrize("path", ["/api/login", "/api/logout"])
def test_password_routes_removed(web, path):
    assert web.client.post(path, headers=HEADERS).status_code == 404


def test_expired_and_invalid_sessions_are_replaced_without_password(web):
    initialize_session(web.client)
    expired = web.client.cookies.get("demo_session")
    with web.app.state.store.connect() as db:
        db.execute("UPDATE sessions SET expires=?", (time.time() - 1,))
    response = initialize_session(web.client)
    assert "set-cookie" in response.headers
    assert web.client.cookies.get("demo_session") != expired
    assert web.app.state.store.session_owner(expired) is None
    web.client.cookies.clear()
    web.client.cookies.set("demo_session", "invalid", domain="127.0.0.1", path="/")
    initialize_session(web.client)
    assert web.client.cookies.get("demo_session") != "invalid"


def test_customer_page_is_upload_focused_and_all_script_elements_exist(web):
    page = web.client.get("/").text
    assert "Upload and evaluate" in page
    assert "Decision Model Comparison" in page
    assert "Prepare your test cases" not in page
    assert 'id="question"' not in page
    assert 'id="dataset-summary"' not in page
    assert 'id="summary-rows"' in page
    assert "Accuracy" in page
    assert "Brier" in page
    assert "Expected intent" in page
    assert "Noul" in page
    assert ">Evaluate</button>" in page
    assert "password" not in page.lower()
    assert "Unlock demo" not in page
    assert "Lock demo" not in page
    script = (ROOT / "comparison.js").read_text(encoding="utf-8")
    script_ids = set(re.findall(r'\$\("([^"]+)"\)', script))
    page_ids = set(re.findall(r'id="([^"]+)"', page))
    assert script_ids <= page_ids
    assert "toFixed" in script
    assert "numeric(probability, 2)" in script
    assert "judged_correct" in script
    assert "evaluator_correct" in script
    assert "/api/login" not in script
    assert "/api/logout" not in script
    assert 'await api("/api/config")' in script
    assert 'sessionStorage.getItem("evaluation-job")' in script
    assert 'sessionStorage.setItem("evaluation-job", result.id)' in script
    assert script.count('sessionStorage.removeItem("evaluation-job")') == 3


def test_status_and_json_restore_metrics_without_rounding_model_values(web):
    initialize_session(web.client)
    owner = web.app.state.store.session_owner(web.client.cookies.get("demo_session"))
    test_case = EvaluationCase(
        id="P",
        message="Cancel the request",
        predicted_intent="cancel_request",
        expected_intent="cancel_request",
    )
    report = report_for([test_case])
    probability = 0.9975273749375825
    report["results"] = [
        {
            "provider": "microsoft",
            "case_id": "P",
            "requested_model": "test-deployment",
            "status": "ok",
            "response_model": "actual-model",
            "correctness_probability": probability,
            "judged_correct": True,
            "evaluator_correct": True,
            "usage": None,
            "latency_ms": 12.3,
        }
    ]
    identifier = web.app.state.store.reserve(owner, report)
    web.app.state.store.update(identifier, "completed", report)
    status = web.client.get(f"/api/jobs/{identifier}").json()["report"]
    download = web.client.get(f"/api/jobs/{identifier}/download?format=json").json()
    assert status == download
    assert status["results"][0]["answers"]["intent_is_correct"]["noul"] == probability
    assert "summary" in status
    assert status["cases"][0]["expected_intent"] == "cancel_request"
    assert status["results"][0]["judged_correct"] is True
    html = web.client.get(f"/api/jobs/{identifier}/download?format=html").text
    assert str(probability) in html
    assert "Accuracy" in html
    assert ">1.00<br>Intent judged correct" in html


def test_cross_site_link_can_open_homepage_and_then_use_demo(web):
    response = web.client.get(
        "/",
        headers={
            "Sec-Fetch-Site": "cross-site",
            "Sec-Fetch-Mode": "navigate",
            "Sec-Fetch-Dest": "document",
        },
    )
    assert response.status_code == 200
    assert "Upload and evaluate" in response.text
    assert response.headers["x-frame-options"] == "DENY"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]
    assert web.client.get("/api/config").status_code == 200


@pytest.mark.parametrize(
    "path,headers",
    [
        ("/", {"Sec-Fetch-Site": "cross-site"}),
        (
            "/",
            {
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "navigate",
                "Sec-Fetch-Dest": "iframe",
            },
        ),
        (
            "/",
            {
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "cors",
                "Sec-Fetch-Dest": "document",
            },
        ),
        (
            "/",
            {
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "navigate",
                "Sec-Fetch-Dest": "document",
                "Origin": "https://evil.invalid",
            },
        ),
        (
            "/api/config",
            {
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "navigate",
                "Sec-Fetch-Dest": "document",
            },
        ),
    ],
)
def test_rejects_cross_site_reads_except_homepage_navigation(web, path, headers):
    initialize_session(web.client)
    response = web.client.get(path, headers=headers)
    assert response.status_code == 403
    assert response.json()["detail"] == "Cross-site requests are not allowed."


@pytest.mark.parametrize(
    "headers",
    [
        {},
        {**HEADERS, "Origin": "https://evil.invalid"},
        {**HEADERS, "Host": "evil.invalid"},
        {**HEADERS, "Sec-Fetch-Site": "cross-site"},
    ],
)
def test_rejects_cross_site_upload(web, headers):
    response = web.client.post("/api/jobs", content=b"workbook", headers=headers)
    assert response.status_code == 403
    assert "set-cookie" not in response.headers


def test_upload_and_owner_isolation(web):
    headers = {
        **HEADERS,
        "Content-Type": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "X-Synthetic-Data": "confirmed",
    }
    content = workbook([["X", "<script>alert(1)</script>", "other", "other"]])
    response = web.client.post("/api/jobs", content=content, headers=headers)
    assert response.status_code == 202
    identifier = response.json()["id"]
    for _ in range(20):
        job = web.client.get(f"/api/jobs/{identifier}").json()
        if job["status"] == "completed":
            break
        time.sleep(0.01)
    assert job["status"] == "completed"
    assert web.client.get(f"/api/jobs/{identifier}/download?format=json").status_code == 200
    page = web.client.get(f"/api/jobs/{identifier}/download?format=html")
    assert "<script>alert(1)</script>" not in page.text
    assert "&lt;script&gt;" in page.text
    web.client.cookies.clear()
    assert web.client.get(f"/api/jobs/{identifier}").status_code == 404
    assert web.client.get(f"/api/jobs/{identifier}/download?format=json").status_code == 404
    assert web.client.get(f"/api/jobs/{identifier}/download?format=html").status_code == 404


def test_upload_requires_consent_and_size_limits(web):
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert web.client.post("/api/jobs", content=b"x", headers=HEADERS).status_code == 415
    headers = {**HEADERS, "Content-Type": content_type}
    assert web.client.post("/api/jobs", content=b"x", headers=headers).status_code == 400
    headers["X-Synthetic-Data"] = "confirmed"
    assert (
        web.client.post("/api/jobs", content=b"x" * (MAX_UPLOAD + 1), headers=headers).status_code
        == 413
    )
    assert web.client.post("/api/jobs", content=b"bad", headers=headers).status_code == 422


def test_unsupported_report_can_be_exported_as_json_only(web):
    initialize_session(web.client)
    owner = web.app.state.store.session_owner(web.client.cookies.get("demo_session"))
    report = {
        "created_at": "earlier",
        "providers": ["jev"],
        "question": {},
        "dataset_sha256": "old",
        "cases": [{"id": "OLD", "message": "Explain this", "expected_intent": "get_information"}],
        "results": [],
        "summary": {},
    }
    identifier = web.app.state.store.reserve(owner, report)
    web.app.state.store.update(identifier, "completed", report)
    response = web.client.get(f"/api/jobs/{identifier}/download?format=html")
    assert response.status_code == 409
    assert "Unsupported report format" in response.json()["detail"]
    assert web.client.get(f"/api/jobs/{identifier}/download?format=json").json() == report


def test_store_global_quota_and_restart(tmp_path):
    store = Store(tmp_path)
    owner = store.session_owner(store.new_session())
    case = EvaluationCase(
        id="X", message="hello", predicted_intent="other", expected_intent="other"
    )
    report = report_for([case])
    identifier = store.reserve(owner, report)
    with pytest.raises(HTTPException) as error:
        store.reserve(owner, report)
    assert error.value.status_code == 409
    restarted = Store(tmp_path)
    assert restarted.get(identifier, owner)["status"] == "failed"
    day = datetime_utc_day()
    with restarted.connect() as db:
        db.execute("UPDATE quota SET rows=? WHERE day=?", (DAILY_ROWS - 1, day))
    second = restarted.reserve(owner, report)
    restarted.update(second, "completed", report)
    with pytest.raises(HTTPException) as error:
        restarted.reserve(owner, report)
    assert error.value.status_code == 429
    with restarted.connect() as db:
        db.execute("UPDATE jobs SET created=?", (time.time() - 86401,))
    with pytest.raises(HTTPException) as error:
        restarted.get(identifier, owner)
    assert error.value.status_code == 404


def datetime_utc_day():
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).date().isoformat()


@pytest.mark.parametrize("failure", [None, "denied", "identity"])
def test_real_runner_with_fake_providers(tmp_path, monkeypatch, failure):
    credential = SimpleNamespace(
        get_token=Mock(return_value=AccessToken("fake-token", 9999999999)), close=Mock()
    )
    if failure == "identity":
        credential.get_token.side_effect = ClientAuthenticationError("private diagnostic")
    monkeypatch.setattr("web_demo.ManagedIdentityCredential", lambda: credential)
    requests = []
    real_client = httpx.Client

    def client(**kwargs):
        def handler(request):
            requests.append(request)
            sent = json.loads(request.content)
            assert "expected_intent" not in sent["state"]
            assert sent["state"]["classifier_predicted_intent"] == "other"
            assert sent["questions"]["intent_is_correct"]["type"] == "noul"
            if failure == "denied":
                return httpx.Response(403)
            return httpx.Response(
                200,
                json={
                    "model": "actual-test-version",
                    "answers": {"intent_is_correct": {"type": "noul", "noul": 0.99}},
                    "usage": {"input_tokens": 123, "output_tokens": 9},
                },
            )

        return real_client(transport=httpx.MockTransport(handler), **kwargs)

    monkeypatch.setattr("web_demo.httpx.Client", client)

    async def no_sleep(seconds):
        return None

    monkeypatch.setattr("web_demo.asyncio.sleep", no_sleep)
    settings = WebSettings(ORIGIN, "fake-key", tmp_path, False)
    store = Store(tmp_path)
    owner = store.session_owner(store.new_session())
    report = report_for(
        [EvaluationCase(id="X", message="hi", predicted_intent="other", expected_intent="other")]
    )
    identifier = store.reserve(owner, report)
    runner = WebRunner(settings, AzureSettings.from_env(), store)
    asyncio.run(runner.run(identifier, report))
    job = store.get(identifier, owner)
    if failure:
        assert job["status"] == "failed"
        assert "private diagnostic" not in json.dumps(job)
        assert len(requests) <= 1
    else:
        assert job["status"] == "completed"
        assert len(requests) == 2
        assert job["report"]["summary"]["jev"]["evaluator_accuracy"] == 1
        assert requests[0].headers["Authorization"] == "Bearer fake-token"
        assert requests[1].headers["Authorization"] == "Bearer fake-key"
        assert "fake-key" not in json.dumps(job)
