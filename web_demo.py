"""Excel comparison demo with anonymous sessions for one Azure App Service instance."""

import asyncio
import hashlib
import io
import json
import logging
import os
import secrets
import sqlite3
import time
import zipfile
from collections.abc import Iterator
from contextlib import asynccontextmanager, contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
from xml.etree.ElementTree import ParseError

import httpx
from azure.core.exceptions import AzureError
from azure.identity import ManagedIdentityCredential
from defusedxml.common import DefusedXmlException
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from intent_evaluator import (
    COLUMNS,
    INTENTS,
    MODE,
    PRICE_PER_MILLION,
    QUESTION,
    ROOT,
    THRESHOLD,
    EvaluationCase,
    evaluate,
    evaluation_report,
    render_report,
    summarize,
)

logger = logging.getLogger(__name__)
SCOPE = "https://cognitiveservices.azure.com/.default"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD = 262144
MAX_ROWS = 100
DAILY_ROWS = 1000
SESSION_SECONDS = 7200
RETENTION_SECONDS = 86400


@dataclass(frozen=True)
class AzureSettings:
    endpoint: str
    deployment: str

    def __post_init__(self) -> None:
        endpoint = urlsplit(self.endpoint)
        if (
            endpoint.scheme != "https"
            or not endpoint.hostname
            or not endpoint.hostname.endswith(
                (".services.ai.azure.com", ".cognitiveservices.azure.com")
            )
            or endpoint.port not in (None, 443)
            or endpoint.username
            or endpoint.password
            or endpoint.path not in ("", "/")
            or endpoint.query
            or endpoint.fragment
        ):
            raise ValueError("AZURE_ENDPOINT must be a root HTTPS Foundry resource endpoint.")
        if not self.deployment.strip():
            raise ValueError("AZURE_DEPLOYMENT cannot be empty.")

    @classmethod
    def from_env(cls) -> "AzureSettings":
        return cls(
            endpoint=os.environ["AZURE_ENDPOINT"],
            deployment=os.environ["AZURE_DEPLOYMENT"],
        )


@dataclass(frozen=True)
class WebSettings:
    origin: str
    jev_key: str
    data_dir: Path
    secure_cookies: bool = True
    jev_model: str = "jev-1.14.0"

    def __post_init__(self) -> None:
        parsed = urlsplit(self.origin)
        if (
            parsed.scheme not in ("https", "http")
            or not parsed.netloc
            or parsed.path
            or parsed.query
            or parsed.fragment
            or parsed.username
            or parsed.password
        ):
            raise ValueError(
                "PUBLIC_ORIGIN must be a root HTTP(S) origin without a trailing slash."
            )
        if not self.secure_cookies and parsed.hostname not in ("127.0.0.1", "localhost"):
            raise ValueError("Insecure cookies are only permitted on loopback.")
        if self.secure_cookies and parsed.scheme != "https":
            raise ValueError("Hosted demo requires HTTPS.")
        if not self.jev_key.strip():
            raise ValueError("Configure JEV_API_KEY.")

    @classmethod
    def from_env(cls) -> "WebSettings":
        return cls(
            origin=os.environ["PUBLIC_ORIGIN"],
            jev_key=os.environ["JEV_API_KEY"],
            data_dir=Path(os.environ.get("DEMO_DATA_DIR", "/home/intent-demo")),
        )


def parse_excel(content: bytes) -> list[EvaluationCase]:
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError("Upload must be a nonempty .xlsx file no larger than 256 KiB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if (
                len(entries) > 128
                or sum(entry.file_size for entry in entries) > 2_000_000
                or len(set(names)) != len(names)
                or any(entry.flag_bits & 1 for entry in entries)
                or any("vba" in name.lower() or "externallinks/" in name.lower() for name in names)
            ):
                raise ValueError(
                    "Workbook is too complex, encrypted, or contains external content."
                )
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("Use exactly one worksheet.")
            sheet = workbook.worksheets[0]
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise ValueError("At most 100 data rows are allowed.")
            if sheet.max_column and sheet.max_column != len(COLUMNS):
                raise ValueError(
                    "Use the evaluator template with id, message, predicted_intent, "
                    "expected_intent."
                )
            rows = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=len(COLUMNS) + 1)
            headers = next(rows)
            if [cell.value for cell in headers[: len(COLUMNS)]] != COLUMNS:
                raise ValueError(
                    "Headers must be id, message, predicted_intent, expected_intent, "
                    "in that order. Download the evaluator template."
                )
            cases = []
            for number, cells in enumerate(rows, start=2):
                if all(cell.value is None for cell in cells):
                    continue
                if len(cases) >= MAX_ROWS or cells[len(COLUMNS)].value is not None:
                    raise ValueError("At most 100 rows and four columns are allowed.")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Row {number}: formulas are not permitted.")
                values = [cell.value for cell in cells[: len(COLUMNS)]]
                if not all(isinstance(value, str) and value.strip() for value in values):
                    raise ValueError(f"Row {number}: all four cells must be nonempty text.")
                identifier, message, predicted, label = [value.strip() for value in values]
                if len(identifier) > 50 or len(message) > 1000:
                    raise ValueError(
                        f"Row {number}: ID exceeds 50 or message exceeds 1,000 characters."
                    )
                if label not in INTENTS or predicted not in INTENTS:
                    raise ValueError(
                        f"Row {number}: predicted_intent and expected_intent "
                        "must be supported intents."
                    )
                cases.append(
                    EvaluationCase(
                        id=identifier,
                        message=message,
                        predicted_intent=predicted,
                        expected_intent=label,
                    )
                )
            if not cases or len({case.id for case in cases}) != len(cases):
                raise ValueError("Include 1-100 cases with unique IDs.")
            return cases
        finally:
            workbook.close()
    except (
        zipfile.BadZipFile,
        InvalidFileException,
        KeyError,
        ParseError,
        DefusedXmlException,
        EOFError,
        ValidationError,
        RuntimeError,
    ) as error:
        raise ValueError("Invalid or unsupported .xlsx workbook.") from error


def excel_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Intent evaluator cases"
    sheet.append(COLUMNS)
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 100
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 30
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class Store:
    def __init__(self, directory: Path) -> None:
        directory.mkdir(parents=True, exist_ok=True)
        self.path = directory / "demo.sqlite3"
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS sessions (id TEXT PRIMARY KEY, expires REAL NOT NULL);
                CREATE TABLE IF NOT EXISTS quota (day TEXT PRIMARY KEY, rows INTEGER NOT NULL);
                CREATE TABLE IF NOT EXISTS jobs (
                    id TEXT PRIMARY KEY, owner TEXT NOT NULL, created REAL NOT NULL,
                    status TEXT NOT NULL, report TEXT NOT NULL, error TEXT
                );
                UPDATE jobs SET status='failed', error='Server restarted; please start a new run.'
                    WHERE status IN ('queued','running');
            """)
        self.cleanup()

    @contextmanager
    def connect(self) -> Iterator[sqlite3.Connection]:
        db = sqlite3.connect(self.path, timeout=10)
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def cleanup(self) -> None:
        with self.connect() as db:
            db.execute("DELETE FROM sessions WHERE expires < ?", (time.time(),))
            db.execute("DELETE FROM jobs WHERE created < ?", (time.time() - RETENTION_SECONDS,))

    def session_owner(self, token: str) -> str | None:
        identifier = hashlib.sha256(token.encode()).hexdigest()
        with self.connect() as db:
            row = db.execute(
                "SELECT id FROM sessions WHERE id=? AND expires>?", (identifier, time.time())
            ).fetchone()
        return identifier if row else None

    def new_session(self) -> str:
        token = secrets.token_urlsafe(32)
        with self.connect() as db:
            db.execute(
                "INSERT INTO sessions VALUES (?,?)",
                (hashlib.sha256(token.encode()).hexdigest(), time.time() + SESSION_SECONDS),
            )
        return token

    def reserve(self, owner: str, report: dict) -> str:
        self.cleanup()
        day = datetime.now(timezone.utc).date().isoformat()
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            active = db.execute(
                "SELECT COUNT(*) FROM jobs WHERE status IN ('queued','running')"
            ).fetchone()[0]
            if active:
                raise HTTPException(409, "Another evaluation is running. Please wait.")
            row = db.execute("SELECT rows FROM quota WHERE day=?", (day,)).fetchone()
            used = row[0] if row else 0
            count = len(report["cases"])
            if used + count > DAILY_ROWS:
                raise HTTPException(429, "The demo's 1,000-row daily allowance has been reached.")
            db.execute(
                "INSERT INTO quota VALUES (?,?) ON CONFLICT(day) DO UPDATE SET rows=excluded.rows",
                (day, used + count),
            )
            identifier = secrets.token_urlsafe(24)
            db.execute(
                "INSERT INTO jobs VALUES (?,?,?,?,?,NULL)",
                (identifier, owner, time.time(), "queued", json.dumps(report)),
            )
        return identifier

    def update(self, identifier: str, status: str, report: dict, error: str | None = None) -> None:
        with self.connect() as db:
            db.execute(
                "UPDATE jobs SET status=?,report=?,error=? WHERE id=?",
                (status, json.dumps(report, allow_nan=False), error, identifier),
            )

    def get(self, identifier: str, owner: str) -> dict:
        self.cleanup()
        with self.connect() as db:
            row = db.execute(
                "SELECT * FROM jobs WHERE id=? AND owner=?", (identifier, owner)
            ).fetchone()
        if not row:
            raise HTTPException(404, "Evaluation not found or no longer retained.")
        return {
            "id": row["id"],
            "status": row["status"],
            "error": row["error"],
            "report": json.loads(row["report"]),
        }


class WebRunner:
    def __init__(self, settings: WebSettings, azure: AzureSettings, store: Store) -> None:
        self.settings = settings
        self.azure = azure
        self.store = store
        self.credential = ManagedIdentityCredential()

    async def run(self, identifier: str, report: dict) -> None:
        cases = [EvaluationCase.model_validate(case) for case in report["cases"]]
        self.store.update(identifier, "running", report)
        failed = False
        try:
            with (
                httpx.Client(timeout=30, follow_redirects=False) as microsoft,
                httpx.Client(
                    headers={"Authorization": f"Bearer {self.settings.jev_key}"},
                    timeout=30,
                    follow_redirects=False,
                ) as jev,
            ):
                for index, case in enumerate(cases):
                    order = ("microsoft", "jev") if index % 2 == 0 else ("jev", "microsoft")
                    for provider in order:
                        if provider == "microsoft":
                            token = await asyncio.to_thread(self.credential.get_token, SCOPE)
                            microsoft.headers["Authorization"] = f"Bearer {token.token}"
                            client = microsoft
                            endpoint = (
                                f"{self.azure.endpoint.rstrip('/')}"
                                "/providers/microsoft/v1/systemone"
                            )
                            model = self.azure.deployment
                        else:
                            client, endpoint, model = (
                                jev,
                                "https://api.typesafe.ai/v1/systemone",
                                self.settings.jev_model,
                            )
                        result = await asyncio.to_thread(
                            evaluate, client, endpoint, model, case, provider
                        )
                        report["results"].append(result)
                        report["summary"] = {
                            p: summarize(cases, report["results"], p) for p in report["providers"]
                        }
                        self.store.update(identifier, "running", report)
                        if result["status"] != "ok":
                            failed = True
                            if result["error"] in {"HTTP 401", "HTTP 403", "HTTP 429"}:
                                self.store.update(
                                    identifier,
                                    "failed",
                                    report,
                                    f"{provider} returned {result['error']}. "
                                    "Evaluation stopped; contact the demo owner.",
                                )
                                return
                    await asyncio.sleep(1.1)
            self.store.update(
                identifier,
                "failed" if failed else "completed",
                report,
                "Some model calls failed. Results are incomplete." if failed else None,
            )
        except AzureError as error:
            logger.error(
                "Managed identity inference authentication failed (%s).", type(error).__name__
            )
            self.store.update(
                identifier,
                "failed",
                report,
                "Microsoft inference authentication failed. Contact the demo owner.",
            )
        except asyncio.CancelledError:
            self.store.update(identifier, "failed", report, "Server stopped during evaluation.")
            raise
        except Exception:
            logger.exception("Evaluation failed unexpectedly.")
            self.store.update(
                identifier,
                "failed",
                report,
                "Evaluation failed unexpectedly. Contact the demo owner.",
            )
            raise


async def bounded_body(request: Request, maximum: int) -> bytes:
    body = bytearray()
    async for chunk in request.stream():
        body.extend(chunk)
        if len(body) > maximum:
            raise HTTPException(413, f"Request exceeds the {maximum}-byte limit.")
    return bytes(body)


def create_app(settings: WebSettings | None = None, *, runner_factory=WebRunner) -> FastAPI:
    settings = settings or WebSettings.from_env()
    azure = AzureSettings.from_env()

    @asynccontextmanager
    async def lifespan(app: FastAPI):
        app.state.store = Store(settings.data_dir)
        app.state.runner = runner_factory(settings, azure, app.state.store)
        app.state.tasks = set()

        async def expire_data() -> None:
            while True:
                await asyncio.sleep(60)
                app.state.store.cleanup()

        cleanup_task = asyncio.create_task(expire_data())
        try:
            yield
        finally:
            cleanup_task.cancel()
            for task in app.state.tasks:
                task.cancel()
            await asyncio.gather(cleanup_task, *app.state.tasks, return_exceptions=True)
            app.state.runner.credential.close()

    app = FastAPI(lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)
    expected_host = urlsplit(settings.origin).netloc

    @app.middleware("http")
    async def protect(request: Request, call_next):
        if request.headers.get("host") != expected_host:
            return JSONResponse({"detail": "Unexpected host."}, status_code=403)
        origin = request.headers.get("origin")
        homepage_navigation = (
            request.method == "GET"
            and request.url.path == "/"
            and origin is None
            and request.headers.get("sec-fetch-mode") == "navigate"
            and request.headers.get("sec-fetch-dest") == "document"
        )
        if (origin is not None and origin != settings.origin) or (
            request.headers.get("sec-fetch-site") == "cross-site" and not homepage_navigation
        ):
            return JSONResponse({"detail": "Cross-site requests are not allowed."}, status_code=403)
        if request.method == "POST" and (
            origin != settings.origin or request.headers.get("x-demo-request") != "1"
        ):
            return JSONResponse({"detail": "Same-origin demo request required."}, status_code=403)
        session_token = None
        if request.url.path.startswith("/api/"):
            owner = app.state.store.session_owner(request.cookies.get("demo_session", ""))
            if owner is None:
                session_token = app.state.store.new_session()
                owner = hashlib.sha256(session_token.encode()).hexdigest()
            request.state.owner = owner
        response = await call_next(request)
        if session_token is not None:
            response.set_cookie(
                "demo_session",
                session_token,
                httponly=True,
                secure=settings.secure_cookies,
                samesite="strict",
                max_age=SESSION_SECONDS,
            )
        response.headers.update(
            {
                "Cache-Control": "no-store",
                "X-Content-Type-Options": "nosniff",
                "Referrer-Policy": "no-referrer",
                "X-Frame-Options": "DENY",
                "Content-Security-Policy": (
                    "default-src 'self'; script-src 'self'; style-src 'self'; "
                    "connect-src 'self'; frame-ancestors 'none'; "
                    "base-uri 'none'; form-action 'self'"
                ),
            }
        )
        if settings.secure_cookies:
            response.headers["Strict-Transport-Security"] = "max-age=31536000"
        return response

    @app.get("/")
    async def index():
        return FileResponse(ROOT / "comparison.html")

    @app.get("/comparison.js")
    async def javascript():
        return FileResponse(ROOT / "comparison.js", media_type="text/javascript")

    @app.get("/comparison.css")
    async def stylesheet():
        return FileResponse(ROOT / "comparison.css", media_type="text/css")

    @app.get("/health")
    async def health():
        return {"status": "ok"}

    @app.get("/api/config")
    async def config():
        return {
            "mode": MODE,
            "judgment_threshold": THRESHOLD,
            "intents": INTENTS,
            "question": QUESTION,
            "max_rows": MAX_ROWS,
            "max_upload_bytes": MAX_UPLOAD,
            "daily_rows": DAILY_ROWS,
            "jev_model": settings.jev_model,
            "microsoft_deployment": azure.deployment,
        }

    @app.get("/api/template")
    async def template():
        return Response(
            excel_template(),
            media_type=XLSX_TYPE,
            headers={"Content-Disposition": 'attachment; filename="evaluator-template.xlsx"'},
        )

    @app.post("/api/jobs")
    async def start(request: Request):
        if request.headers.get("content-type") != XLSX_TYPE:
            raise HTTPException(415, "Upload an .xlsx workbook.")
        if request.headers.get("x-synthetic-data") != "confirmed":
            raise HTTPException(400, "Confirm that the workbook contains only synthetic test data.")
        content = await bounded_body(request, MAX_UPLOAD)
        try:
            cases = await asyncio.to_thread(parse_excel, content)
        except ValueError as error:
            raise HTTPException(422, str(error)) from error
        report = {
            "mode": MODE,
            "judgment_threshold": THRESHOLD,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "providers": ["microsoft", "jev"],
            "question": QUESTION,
            "cases": [case.model_dump() for case in cases],
            "dataset_sha256": hashlib.sha256(content).hexdigest(),
            "jev_model": settings.jev_model,
            "azure_deployment": azure.deployment,
            "price_usd_per_million_input_tokens": PRICE_PER_MILLION,
            "results": [],
            "summary": {p: summarize(cases, [], p) for p in ("microsoft", "jev")},
        }
        identifier = app.state.store.reserve(request.state.owner, report)
        task = asyncio.create_task(app.state.runner.run(identifier, report))
        app.state.tasks.add(task)

        def finished(task: asyncio.Task) -> None:
            app.state.tasks.discard(task)
            if not task.cancelled() and task.exception() is not None:
                logger.error("Evaluation task terminated with an unexpected exception.")

        task.add_done_callback(finished)
        return JSONResponse({"id": identifier}, status_code=202)

    @app.get("/api/jobs/{identifier}")
    async def status(identifier: str, request: Request):
        job = app.state.store.get(identifier, request.state.owner)
        if job["report"].get("mode") == MODE:
            job["report"] = evaluation_report(job["report"])
        return job

    @app.get("/api/jobs/{identifier}/download")
    async def download(identifier: str, request: Request, format: str = "json"):
        job = app.state.store.get(identifier, request.state.owner)
        if job["status"] in ("queued", "running"):
            raise HTTPException(409, "Wait until the evaluation finishes.")
        if format == "html":
            if job["report"].get("mode") != MODE:
                raise HTTPException(
                    409, "Unsupported report format. Download JSON or start a new evaluation."
                )
            return HTMLResponse(
                render_report(job["report"]),
                headers={
                    "Content-Disposition": 'attachment; filename="comparison.html"',
                    "Content-Security-Policy": (
                        "default-src 'none'; style-src 'unsafe-inline'; "
                        "frame-ancestors 'none'; base-uri 'none'"
                    ),
                },
            )
        if format != "json":
            raise HTTPException(422, "Choose json or html.")
        return JSONResponse(
            evaluation_report(job["report"])
            if job["report"].get("mode") == MODE
            else job["report"],
            headers={"Content-Disposition": 'attachment; filename="comparison.json"'},
        )

    return app
